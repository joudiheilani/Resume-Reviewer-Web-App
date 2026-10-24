"""Flask web server for the AI Resume Reviewer."""

import io
import os

from flask import Flask, jsonify, render_template, request

import reviewer

# Static files live in public/ so Vercel can serve them from its CDN; Flask serves them locally.
app = Flask(__name__, static_folder="public/static", static_url_path="/static")
app.config["MAX_CONTENT_LENGTH"] = 4 * 1024 * 1024  # 4 MB; Vercel caps request bodies at 4.5 MB

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def extract_text(filename, data):
    """Return plain text from an uploaded PDF, DOCX or text file."""
    ext = os.path.splitext(filename.lower())[1]
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError("Please upload a PDF, DOCX or TXT file.")
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif ext == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    else:
        text = data.decode("utf-8", errors="ignore")
    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not text:
        raise ValueError(
            "No text could be read from that file. If it is a scanned PDF, paste the text instead."
        )
    return text


def _text_field(body, name, label, max_chars, min_chars=1):
    value = body.get(name)
    if not isinstance(value, str) or len(value.strip()) < min_chars:
        raise ValueError(f"{label} is required.")
    return value.strip()[:max_chars]


def _inputs(body):
    resume = _text_field(body, "resume", "Your resume", reviewer.MAX_RESUME_CHARS, min_chars=50)
    job = _text_field(body, "job_description", "The job description", reviewer.MAX_JOB_CHARS, 30)
    role = body.get("role_level")
    reviewer.role_guidance(role)  # raises ValueError for anything but internship / entry_level
    context = body.get("additional_context")
    context = context.strip()[: reviewer.MAX_CONTEXT_CHARS] if isinstance(context, str) else ""
    return resume, job, role, context


@app.errorhandler(ValueError)
def bad_request(exc):
    return jsonify(error=str(exc)), 400


@app.errorhandler(reviewer.AIError)
def ai_unavailable(exc):
    return jsonify(error=str(exc)), 503


@app.errorhandler(413)
def too_large(_exc):
    return jsonify(error="That file is too large. The limit is 4 MB."), 413


@app.get("/")
def index():
    return render_template("index.html", model=reviewer.get_models()[0])


@app.post("/api/extract")
def extract():
    upload = request.files.get("file")
    if not upload or not upload.filename:
        raise ValueError("No file was uploaded.")
    text = extract_text(upload.filename, upload.read())
    return jsonify(text=text[: reviewer.MAX_RESUME_CHARS], filename=upload.filename)


@app.post("/api/review/draft")
def review_draft():
    """Step 1: Persona pattern. The recruiter writes the first review."""
    resume, job, role, context = _inputs(request.get_json(silent=True) or {})
    draft, model = reviewer.first_review(resume, job, role, context)
    return jsonify(draft=draft, model=model)


@app.post("/api/review/refine")
def review_refine():
    """Step 2: Iterative Refinement pattern. A strict editor improves the first review."""
    body = request.get_json(silent=True) or {}
    resume, job, role, context = _inputs(body)
    draft = _text_field(body, "draft", "The first review", 20000)
    feedback, model = reviewer.refine_review(resume, job, draft, role, context)
    return jsonify(feedback=feedback, model=model)


@app.post("/api/chat")
def chat():
    """Follow-up questions about the resume and feedback."""
    body = request.get_json(silent=True) or {}
    resume, job, role, context = _inputs(body)
    feedback = _text_field(body, "feedback", "The feedback", 20000)
    question = _text_field(body, "question", "A question", reviewer.MAX_QUESTION_CHARS)
    history = body.get("history") if isinstance(body.get("history"), list) else []
    answer, model = reviewer.answer_question(
        resume, job, feedback, history, question, role, context
    )
    return jsonify(answer=answer, model=model)


if __name__ == "__main__":
    port = int(os.getenv("PORT", "5000"))
    app.run(host="127.0.0.1", port=port, debug=os.getenv("FLASK_DEBUG") == "1")
